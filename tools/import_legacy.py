"""Conversor local para la estructura inspeccionada de 4 de Abril MEJORADO- 2025.xlsx.
No modifica el original. Usa valores guardados: no recalcula fórmulas de Google Sheets.
"""
import argparse
import datetime as dt
import json
from pathlib import Path
import re
import uuid
import openpyxl

def uid(): return str(uuid.uuid4())
def text(v):
    if v is None or isinstance(v,str) and v.startswith(('#N/A','#REF!','#VALUE!')): return ''
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()
def date(v):
    if isinstance(v,(dt.datetime,dt.date)): return v.date().isoformat() if isinstance(v,dt.datetime) else v.isoformat()
    for fmt in ('%d/%m/%Y','%Y-%m-%d'):
        try: return dt.datetime.strptime(text(v),fmt).date().isoformat()
        except ValueError: pass
    return ''
def doc(v): return re.sub(r'[.\s-]','',text(v))

def convert(source):
    w=openpyxl.load_workbook(source,read_only=True,data_only=True)
    s=w['Colectivo']
    rows=list(s.iter_rows(min_row=1,max_row=200,max_col=42,values_only=True))
    def value(address): return s[address].value
    if text(value('Z5'))!='NOMBRE' or text(value('AA5'))!='APELLIDO':
        raise ValueError('El archivo no coincide con el esquema documentado: Z5/AA5.')
    # Las posiciones se toman del plano visible, no de una plantilla inventada.
    floors=[]
    for name,columns,start,end in [('Piso superior',[3,4,6],8,28),('Piso inferior',[8,9,11],14,20)]:
        seats=[]
        for r in range(start,end+1,2):
            for c,col in zip(columns,[0,1,3]):
                label=text(rows[r-1][c-1])
                if label.isdigit(): seats.append(dict(id=uid(),label=label,row=(r-start)//2,col=col))
        floors.append(dict(id=uid(),name=name,rows=(end-start)//2+2,cols=4,seats=seats))
    t=dict(id=uid(),name=Path(source).stem,origin='Córdoba',destination='Brasil',departure=date(value('Y1')),returnDate=date(value('Y2')),checkIn=date(value('AE1')),checkOut=date(value('AE2')),presentation=text(value('AB1')),departureTime=text(value('AB2')),boardingPlace=text(value('AE3')),coordinator=text(value('AE4')),carrier='',vehicle='',plate='',route='',border='',provider='',floors=floors,groups=[],passengers=[],configured=False)
    manifest=w['Manifiesto']
    for key,cell in [('carrier','C8'),('vehicle','D8'),('plate','E8'),('route','D6'),('border','B6')]:t[key]=text(manifest[cell].value)
    # Habitación y camas pueden estar en celdas combinadas. Propagar únicamente dentro del mismo grupo.
    room_by_doc={}
    roomrows=list(w['Rooming'].iter_rows(min_row=7,max_row=100,max_col=80,values_only=True))
    for offset in range(0,80,10):
        last_group=None; inherited={}
        for r in roomrows:
            g=text(r[offset]); d=doc(r[offset+3])
            if not g or g!=last_group: inherited={}
            last_group=g
            for key,idx in [('meal',6),('roomType',7),('beds',8),('room',9)]:
                v=text(r[offset+idx])
                if v: inherited[key]=v
            if d: room_by_doc[d]=dict(inherited)
    groups={};warnings=[];colors=['#087f73','#4964a5','#945b9c','#b15d34','#487747','#936e24','#297b9b','#a64962']
    for rownum,r in enumerate(rows,1):
        if rownum not in (6,7) and rownum<9: continue
        if not text(r[25]) or not text(r[26]):continue
        crew=rownum in (6,7)
        key=text(r[22]) if not crew else ''
        if key and key not in groups:
            g=dict(id=uid(),name='Grupo '+key+' · '+text(r[24]),color=colors[len(groups)%len(colors)])
            groups[key]=g;t['groups'].append(g)
        p=dict(id=uid(),firstName=text(r[25]),lastName=text(r[26]),document=text(r[27]),documentType='DNI',nationality=text(r[28]),residence=text(r[33]),occupation=text(r[29]),sex=text(r[30]),birthDate=date(r[31]),phone=text(r[38]),groupId=groups[key]['id'] if key else '',seatId='',noSeat=crew or text(r[14]).lower() in ('no','false'),origin='',destination=text(r[34] or r[17]),boarding=text(r[39] or r[41]),hotel=text(r[18] or r[15]),roomType=text(r[19] or r[36]),room='',beds='',meal=text(r[35]),notes=text(r[37]),boarded=False,role='Tripulante' if crew else 'Pasajero')
        if text(r[16]):p['notes']=(p['notes']+'\nOrigen comercial: '+text(r[16])).strip()
        p.update(room_by_doc.get(doc(r[27]),{}))
        label=text(r[20]);seat=next((a for f in floors for a in f['seats'] if a['label']==label),None)
        if seat and not crew:p['seatId']=seat['id'];p['noSeat']=False
        elif label:warnings.append('Fila '+str(rownum)+': butaca sin equivalencia en el plano.')
        t['passengers'].append(p)
    warnings.extend(['Los valores proceden de la copia guardada del Excel; no se recalcularon fórmulas.', 'La columna U no contiene asignaciones en las filas de pasajeros de este archivo: quedan pendientes.', 'El ORIGEN de la columna Q se conservó como origen comercial en observaciones; no es una ciudad.', 'Los dos tripulantes se importaron sin butaca de pasajero; revisar antes de usar.', 'El plano histórico tiene 43 butacas. Confirmarlo en la aplicación después de revisarlo.', 'No se importan activadores, macros, enlaces de Drive, imágenes ni propiedades de scripts.'])
    w.close()
    return {'version':1,'trips':[t]},warnings

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source');parser.add_argument('--output',default='.local-data/abril-2025.json')
    args=parser.parse_args();data,warnings=convert(args.source);out=Path(args.output)
    out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    out.with_suffix('.notas.txt').write_text('\n'.join(warnings),encoding='utf-8')
    print(json.dumps({'passengers':len(data['trips'][0]['passengers']),'groups':len(data['trips'][0]['groups']),'seats':sum(len(f['seats']) for f in data['trips'][0]['floors']),'output':str(out)},ensure_ascii=True))
